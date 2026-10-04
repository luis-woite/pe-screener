"""
Phase 4 — Memo generation.

Assembles everything from Phases 1-3 (snapshot, metrics, LLM synthesis)
into a clean, formatted Word document — the actual deliverable a
screening tool should produce.
"""

import os
from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

ACCENT_COLOR = RGBColor(0x1F, 0x38, 0x64)  # dark navy, matches the project plan doc
LIGHT_GREY = RGBColor(0x66, 0x66, 0x66)


def _add_heading(doc: Document, text: str, size: int = 16):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = ACCENT_COLOR
    p.space_after = Pt(6)
    return p


def _add_metric_table(doc: Document, metrics: dict):
    trend = metrics.get("trend_table")
    leverage = metrics.get("leverage", {})
    multiples = metrics.get("multiples", {})

    def fmt(value, suffix=""):
        if value is None:
            return "N/A"
        try:
            import math

            if isinstance(value, float) and math.isnan(value):
                return "N/A"
        except TypeError:
            pass
        return f"{value}{suffix}"

    latest = None
    if trend is not None and not trend.empty:
        latest = trend[trend.columns[0]]

    rows = [
        ("Revenue growth (YoY)", fmt(latest["revenue_growth_pct"], "%") if latest is not None else "N/A"),
        ("Gross margin", fmt(latest["gross_margin_pct"], "%") if latest is not None else "N/A"),
        ("EBITDA margin", fmt(latest["ebitda_margin_pct"], "%") if latest is not None else "N/A"),
        ("Net margin", fmt(latest["net_margin_pct"], "%") if latest is not None else "N/A"),
        ("Net Debt / EBITDA", fmt(leverage.get("net_debt_to_ebitda"), "x")),
        ("Debt / Equity", fmt(leverage.get("debt_to_equity"), "x")),
        ("EV / EBITDA", fmt(multiples.get("ev_to_ebitda"), "x")),
        ("P/E (proxy)", fmt(multiples.get("trailing_pe_proxy"), "x")),
    ]

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Light Grid Accent 1"

    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    for label, value in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = value

    return table


def _add_bullet_list(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def generate_memo(
    ticker: str,
    snapshot: dict,
    metrics: dict,
    synthesis: dict,
    output_dir: str = "output",
    model_signal: dict = None,
) -> str:
    """
    Builds the screening memo .docx and saves it to output_dir.
    Returns the path to the saved file.
    """
    os.makedirs(output_dir, exist_ok=True)

    company_name = snapshot.get("name") or ticker
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # --- Title block ---
    title = doc.add_paragraph()
    title_run = title.add_run(f"Preliminary Screening Memo — {company_name}")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = ACCENT_COLOR

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run(
        f"{ticker} | {snapshot.get('sector') or 'N/A'} — {snapshot.get('industry') or 'N/A'} | "
        f"Generated {date.today().isoformat()}"
    )
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = LIGHT_GREY

    doc.add_paragraph()  # spacer

    # --- Executive summary ---
    _add_heading(doc, "Executive Summary")
    if "error" in synthesis:
        doc.add_paragraph(
            f"LLM synthesis was unavailable for this run ({synthesis['error']}). "
            "This memo contains only the computed financial metrics below."
        )
    else:
        doc.add_paragraph(synthesis.get("summary", "No summary available."))

    doc.add_paragraph()

    # --- Key metrics ---
    _add_heading(doc, "Key Financial Metrics")
    _add_metric_table(doc, metrics)
    doc.add_paragraph()

    # --- Risk flags ---
    if "error" not in synthesis and synthesis.get("risk_flags"):
        _add_heading(doc, "Risk Flags")
        _add_bullet_list(doc, synthesis["risk_flags"])
        doc.add_paragraph()

    # --- Key themes ---
    if "error" not in synthesis and synthesis.get("key_themes"):
        _add_heading(doc, "Key Themes")
        _add_bullet_list(doc, synthesis["key_themes"])
        doc.add_paragraph()

    # --- Model signal (Phase 6, optional) ---
    if model_signal and "error" not in model_signal:
        _add_heading(doc, "Model Signal (ML)")
        prob = model_signal.get("probability_expansion")
        model_name = model_signal.get("model_name", "model")
        direction = "margin expansion" if prob is not None and prob >= 0.5 else "margin contraction"
        confidence = prob if (prob is not None and prob >= 0.5) else (1 - prob if prob is not None else None)
        if confidence is not None:
            doc.add_paragraph(
                f"A {model_name} classifier, trained on historical financials across ~70 companies, "
                f"estimates a {confidence:.0%} probability of {direction} next year, based on this "
                f"company's current margin, growth, and leverage profile."
            )
        doc.add_paragraph()

    # --- Disclaimer footer ---
    disclaimer = doc.add_paragraph()
    disclaimer_run = disclaimer.add_run(
        "This memo was generated automatically from public financial data, SEC filings, and news sources, "
        "synthesised with an LLM. It is intended as a first-pass screening aid only and is not a substitute "
        "for full due diligence. Figures and commentary should be independently verified before use in any "
        "investment decision."
    )
    disclaimer_run.italic = True
    disclaimer_run.font.size = Pt(8)
    disclaimer_run.font.color.rgb = LIGHT_GREY

    # Timestamped to the second so reruns never collide with (or get blocked by)
    # a previous file still open in Word, locked by OneDrive, etc.
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    filename = f"{ticker}_screening_memo_{timestamp}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath


if __name__ == "__main__":
    # Manual test: python -m src.memo.generate DE
    import sys

    sys.path.insert(0, ".")
    from src.ingestion.financials import get_company_snapshot, get_financial_statements
    from src.ingestion.filings import get_latest_filings, fetch_filing_text
    from src.ingestion.news import get_recent_news
    from src.metrics.compute import get_full_metrics
    from src.llm.synthesize import synthesize_screening_summary
    from src.ml.predict import predict_margin_trajectory
    from config import OUTPUT_DIR

    ticker = sys.argv[1] if len(sys.argv) > 1 else "DE"

    snap = get_company_snapshot(ticker)
    stmts = get_financial_statements(ticker)
    metrics = get_full_metrics(snap, stmts)
    filings = get_latest_filings(ticker)
    excerpt = fetch_filing_text(filings[0]["primary_doc_url"], max_chars=8000) if filings else ""
    news = get_recent_news(snap.get("name") or ticker)
    synthesis = synthesize_screening_summary(snap, metrics, excerpt, news)
    model_signal = predict_margin_trajectory(snap, metrics)

    path = generate_memo(ticker, snap, metrics, synthesis, output_dir=OUTPUT_DIR, model_signal=model_signal)
    print(f"Memo saved to: {path}")